import io
import json

from docx import Document
import pdfplumber
from pdfminer.pdftypes import resolve1
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from .models import UserProfile, UserRole
from .views import build_docx_from_markdown, build_pdf_from_markdown, parse_markdown_root


REPORT = """# 合成报告标题

## 临床资料

普通中文段落，**粗体示例**、*斜体示例*与***粗斜体示例***。第二行包含 `中文代码`。

参考 [合成文档](https://example.org/guide?a=1&b=2)，~~删除线示例~~。

### 检查结果

1. 首个有序条目
    - 嵌套无序条目\x20\x20
      显式换行后内容
    - 第二个嵌套条目
2. 第二个有序条目

> 引用中的中文提示。
>
> 引用第二段。

| 检查 | 合成结果 |
| --- | --- |
| EEG | **中文脑电检查内容** |
| MRI | 合成影像描述 |

---

```text
中文代码块
if value < 3:
    print("测试")
```
"""


class ReportPDFTests(TestCase):
    def pdf_text(self, data):
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)

    def test_markdown_structure_and_chinese_text_survive_pdf_export(self):
        data = build_pdf_from_markdown(REPORT, title="合成报告标题").getvalue()
        self.assertTrue(data.startswith(b"%PDF-"))
        text = self.pdf_text(data)
        for expected in (
            "临床资料", "检查结果", "粗体示例", "斜体示例", "粗斜体示例",
            "中文代码", "首个有序条目", "嵌套无序条目", "第二个有序条目",
            "引用中的中文提示", "中文脑电检查内容", "合成影像描述", "中文代码块", "第 1 页",
        ):
            with self.subTest(expected=expected):
                self.assertIn(expected, text)
        self.assertEqual(text.count("合成报告标题"), 1)
        self.assertNotIn("**", text)
        self.assertNotIn("```", text)

    def test_long_tables_paginate_with_repeated_headers(self):
        rows = [f"| ROW-{index:03d} | 合成检查描述需要自动换行，提供充足说明文本。 |" for index in range(90)]
        source = "# 分页合成报告\n\n| 编号 | 描述表头 |\n| --- | --- |\n" + "\n".join(rows)
        with pdfplumber.open(build_pdf_from_markdown(source)) as pdf:
            self.assertGreaterEqual(len(pdf.pages), 3)
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            self.assertIn("ROW-000", text)
            self.assertIn("ROW-089", text)
            for index, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                self.assertIn("描述表头", page_text)
                self.assertIn(f"第 {index} 页", page_text)
                for character in page.chars:
                    self.assertGreaterEqual(character["x0"], 45)
                    self.assertLessEqual(character["x1"], page.width - 45)

    def test_a_table_row_taller_than_a_page_can_split(self):
        description = "超长单元格合成内容用于验证跨页分割。" * 220 + "单元格结束标记"
        source = f"# 长单元格\n\n| 项目 | 描述表头 |\n| --- | --- |\n| 测试项 | {description} |"
        with pdfplumber.open(build_pdf_from_markdown(source)) as pdf:
            self.assertGreaterEqual(len(pdf.pages), 2)
            self.assertIn("单元格结束标记", "\n".join(page.extract_text() or "" for page in pdf.pages))
            for page in pdf.pages:
                self.assertIn("描述表头", page.extract_text() or "")

    def test_empty_report_still_produces_a_valid_pdf(self):
        data = build_pdf_from_markdown("").getvalue()
        self.assertIn("患者报告", self.pdf_text(data))

    def test_chinese_regular_bold_and_italic_faces_are_embedded(self):
        with pdfplumber.open(build_pdf_from_markdown(REPORT)) as pdf:
            faces = {character["fontname"] for page in pdf.pages for character in page.chars}
            self.assertGreaterEqual(len(faces), 4)
            embedded = 0
            for reference in resolve1(pdf.pages[0].page_obj.resources["Font"]).values():
                font = resolve1(reference)
                if "FontDescriptor" in font:
                    descriptor = resolve1(font["FontDescriptor"])
                    self.assertIn("FontFile2", descriptor)
                    embedded += 1
            self.assertGreaterEqual(embedded, 4)

    def test_safe_markdown_links_survive_as_pdf_links(self):
        with pdfplumber.open(build_pdf_from_markdown(REPORT)) as pdf:
            links = [link["uri"] for page in pdf.pages for link in page.hyperlinks]
            self.assertIn("https://example.org/guide?a=1&b=2", links)

    def test_list_hard_line_breaks_are_preserved(self):
        source = "- 换行首句  \n  换行次句"
        text = self.pdf_text(build_pdf_from_markdown(source).getvalue())
        self.assertIn("换行首句\n换行次句", text)

    def test_normal_root_and_word_tables_remain_compatible(self):
        root = parse_markdown_root(REPORT)
        self.assertEqual(root[0].tag, "h1")
        self.assertTrue(any(node.tag == "table" for node in root))
        word = Document(build_docx_from_markdown(REPORT))
        self.assertEqual(len(word.tables), 1)
        self.assertEqual(len(word.tables[0].columns), 2)
        self.assertEqual(word.tables[0].rows[1].cells[1].text, "中文脑电检查内容")


class ReportPDFResponseTests(TestCase):
    def test_pdf_response_and_existing_read_permissions(self):
        for role in (UserRole.ADMIN, UserRole.STAFF, UserRole.GUEST):
            with self.subTest(role=role):
                user = get_user_model().objects.create_user(username=f"pdf_{role.lower()}")
                UserProfile.objects.create(user=user, role=role)
                self.client.force_login(user)
                response = self.client.post(
                    reverse("epilepsy:report_export", args=["pdf"]),
                    json.dumps({"markdown": REPORT, "patient_name": "合成示例"}),
                    content_type="application/json",
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "application/pdf")
                self.assertIn("attachment", response["Content-Disposition"])
                self.assertTrue(b"".join(response.streaming_content).startswith(b"%PDF-"))
                response.close()

    def test_export_requires_login_and_a_profile(self):
        url = reverse("epilepsy:report_export", args=["pdf"])
        self.assertEqual(self.client.post(url).status_code, 302)
        user = get_user_model().objects.create_user(username="pdf_unprofiled")
        self.client.force_login(user)
        self.assertEqual(self.client.post(url).status_code, 403)
