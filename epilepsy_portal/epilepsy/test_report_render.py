from docx import Document
from lxml import html

from django.test import SimpleTestCase

from epilepsy.views import (
    build_docx_from_markdown,
    parse_markdown_root,
    render_markdown_to_html,
)


SYNTHETIC_MARKDOWN = """# 合成报告

这是**重点结论**，包含*补充说明*和 `示例代码`。

> 这段内容仅用于排版验收。

- 合成检查一
- 合成检查二

3. 第三步
4. 第四步

| 项目 | 合成结果 |
| --- | --- |
| 检查 A | **正常** |
| 检查 B | 待复核 |

```text
示例代码：不含真实病例。
```
"""


class ReportMarkdownRenderingTests(SimpleTestCase):
    def test_markdown_root_exposes_blocks_without_an_extra_wrapper(self):
        root = parse_markdown_root(SYNTHETIC_MARKDOWN)
        self.assertEqual([node.tag for node in root], [
            "h1", "p", "blockquote", "ul", "ol", "table", "pre",
        ])
        self.assertEqual(root.find("./h1").text_content(), "合成报告")
        cells = root.xpath("./table/tbody/tr[1]/td")
        self.assertEqual([cell.text_content() for cell in cells], ["检查 A", "正常"])

    def test_markdown_preserves_inline_formatting_and_numbering(self):
        root = parse_markdown_root(SYNTHETIC_MARKDOWN)
        self.assertEqual(root.find("./p/strong").text_content(), "重点结论")
        self.assertEqual(root.find("./p/em").text_content(), "补充说明")
        self.assertEqual(root.find("./p/code").text_content(), "示例代码")
        self.assertEqual(root.find("./ol").get("start"), "3")
        self.assertEqual(root.find("./pre/code").text_content().strip(), "示例代码：不含真实病例。")

    def test_active_html_and_unsafe_link_attributes_are_removed(self):
        rendered = render_markdown_to_html("""
<script>synthetic_test()</script>
<iframe src="https://example.test/frame"></iframe>
<svg onload="synthetic_test()"><script>synthetic_test()</script></svg>
<p onclick="synthetic_test()" style="background:url(https://example.test/image)">保留正文</p>
<a href="javascript:synthetic_test()" onmouseover="synthetic_test()">危险链接</a>
<img src="x" onerror="synthetic_test()">
""")
        root = html.fragment_fromstring(rendered, create_parent="div")
        self.assertFalse(root.xpath(".//script | .//iframe | .//svg"))
        self.assertIn("保留正文", root.text_content())
        for node in root.iter():
            for attribute, value in node.attrib.items():
                self.assertFalse(attribute.lower().startswith("on"))
                self.assertNotEqual(attribute.lower(), "style")
                if attribute.lower() in {"href", "src"}:
                    self.assertFalse(value.lower().startswith(("javascript:", "data:", "vbscript:")))

    def test_normal_markdown_links_remain_readable(self):
        rendered = render_markdown_to_html(
            "[参考链接](https://example.test/reference?a=1&b=2) 和 **重要结论**。",
        )
        root = html.fragment_fromstring(rendered, create_parent="div")
        link = root.find(".//a")
        self.assertEqual(link.text_content(), "参考链接")
        self.assertEqual(link.get("href"), "https://example.test/reference?a=1&b=2")
        self.assertEqual(root.find(".//strong").text_content(), "重要结论")

    def test_soft_newlines_and_explicit_line_breaks_match_preview(self):
        soft = parse_markdown_root("第一行\n第二行")
        hard = parse_markdown_root("第一行  \n第二行")
        self.assertFalse(soft.xpath("./p/br"))
        self.assertEqual(len(hard.xpath("./p/br")), 1)
        self.assertEqual(hard.find("./p").text_content(), "第一行\n第二行")

    def test_strikethrough_retains_markdown_semantics(self):
        root = parse_markdown_root("原文 ~~删除的结论~~ 保留正文。")
        self.assertEqual(root.find("./p/del").text_content(), "删除的结论")

    def test_word_export_retains_report_heading_and_table_cells(self):
        document = Document(build_docx_from_markdown(SYNTHETIC_MARKDOWN, title="排版验收"))
        self.assertTrue(any(
            paragraph.text == "合成报告" and paragraph.style.name == "Heading 1"
            for paragraph in document.paragraphs
        ))
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual([cell.text for cell in table.rows[0].cells], ["项目", "合成结果"])
        self.assertEqual([cell.text for cell in table.rows[1].cells], ["检查 A", "正常"])
